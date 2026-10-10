from pptx import Presentation
import sys

def extract_text(ppt_path):
    prs = Presentation(ppt_path)
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                print(shape.text)
        print("\n")
        
extract_text(r"C:\Users\WELCOME\Downloads\ppt_2311981626.pptx")
